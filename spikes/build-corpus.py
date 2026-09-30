"""Builds spikes/corpus/ from the raw downloads in spikes/corpus/_raw/.

Throwaway spike code (excluded from lint and check). Needs: pymupdf, python-docx, pillow.
Provenance and licenses are documented in spikes/corpus/README.md.
"""
import io
import random
from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFilter

ROOT = Path(__file__).parent / "corpus"
RAW = ROOT / "_raw"


def excerpt(src: str, dst: str, first: int, last: int) -> None:
    """Copy pages first..last (1-based, inclusive) into a new text PDF."""
    out = pymupdf.open()
    out.insert_pdf(pymupdf.open(RAW / src), from_page=first - 1, to_page=last - 1)
    out.save(ROOT / dst)


def make_scan(src: str, dst: str, first: int, last: int) -> None:
    """Rasterize pages to an image-only PDF with slight skew, noise and JPEG artifacts."""
    random.seed(7)
    source = pymupdf.open(RAW / src)
    out = pymupdf.open()
    for number in range(first - 1, last):
        pix = source[number].get_pixmap(dpi=150)
        img = Image.open(io.BytesIO(pix.tobytes("png"))).convert("L")
        img = img.rotate(1.2, expand=True, fillcolor=255).filter(ImageFilter.GaussianBlur(0.6))
        pixels = img.load()
        for _ in range(img.width * img.height // 40):
            x, y = random.randrange(img.width), random.randrange(img.height)
            pixels[x, y] = max(0, min(255, pixels[x, y] + random.randint(-40, 40)))
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=55)
        page = out.new_page(width=img.width * 72 / 150, height=img.height * 72 / 150)
        page.insert_image(page.rect, stream=buf.getvalue())
    out.save(ROOT / dst)


def make_docx(dst: str) -> None:
    """Synthetic status report: headings, a table and an embedded chart image (known facts)."""
    doc = Document()
    doc.add_heading("Projekt Nordlicht: Statusbericht Q3 2025", level=0)
    doc.add_heading("1. Zusammenfassung", level=1)
    doc.add_paragraph(
        "Das Projekt Nordlicht wird von Dr. Katharina Brandt geleitet. Das Gesamtbudget beträgt "
        "1,25 Mio. Euro, davon wurden bis Ende September 780.000 Euro verbraucht. Der Go-live "
        "ist für den 14. Februar 2026 geplant."
    )
    doc.add_heading("2. Kennzahlen", level=1)
    rows = [("Quartal", "Umsatz (T€)", "Mitarbeitende"), ("Q1 2025", "312", "18"),
            ("Q2 2025", "398", "21"), ("Q3 2025", "455", "24")]
    table = doc.add_table(rows=len(rows), cols=3)
    table.style = "Table Grid"
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    doc.add_heading("3. Risiken", level=1)
    doc.add_paragraph(
        "Das größte Risiko ist die Verzögerung der Lieferung des Sensormoduls durch die Firma "
        "Helmholz Systems. Ein Ausweichlieferant wurde nicht identifiziert."
    )
    chart = Image.new("RGB", (480, 240), "white")
    draw = ImageDraw.Draw(chart)
    for i, height in enumerate((100, 130, 160)):
        draw.rectangle([60 + i * 130, 220 - height, 140 + i * 130, 220], fill="steelblue")
    buf = io.BytesIO()
    chart.save(buf, "PNG")
    doc.add_heading("4. Umsatzentwicklung", level=1)
    doc.add_picture(io.BytesIO(buf.getvalue()), width=Inches(4))
    doc.save(ROOT / dst)


def main() -> None:
    excerpt("dpr.pdf", "01-dpr-paper-en.pdf", 1, 9)
    excerpt("arbeitsmarkt.pdf", "02-destatis-arbeitsmarkt-de.pdf", 1, 10)
    make_docx("03-projekt-nordlicht-de.docx")
    excerpt("gg.pdf", "04-grundgesetz-auszug-de.pdf", 1, 6)
    make_scan("nist.pdf", "05-nist-ai-rmf-scan-en.pdf", 5, 7)
    text = "".join(page.get_text() for page in list(pymupdf.open(RAW / "bdsg.pdf"))[4:9])
    (ROOT / "06-bdsg-auszug-de.txt").write_text(text, encoding="utf-8")
    (ROOT / "07-wikipedia-rag-en.html").write_bytes((RAW / "rag.html").read_bytes())


main()
