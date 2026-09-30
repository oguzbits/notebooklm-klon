"""Writes the reference text of every corpus file to spikes/corpus/_text/<file>.txt.

Throwaway spike code (excluded from lint and check). Needs: pymupdf, python-docx, beautifulsoup4.
The golden-question test checks every anchor against these files. For the image-only scan the
text of the original pages in `_raw/` is the reference (no OCR).
"""
from pathlib import Path

import pymupdf
from bs4 import BeautifulSoup
from docx import Document

ROOT = Path(__file__).parent / "corpus"
RAW = ROOT / "_raw"
OUT = ROOT / "_text"


def pdf_text(path: Path, first: int = 1, last: int | None = None) -> str:
    doc = pymupdf.open(path)
    return "\n".join(doc[i].get_text() for i in range(first - 1, last or len(doc)))


def docx_text(path: Path) -> str:
    doc = Document(path)
    cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join([paragraph.text for paragraph in doc.paragraphs] + cells)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    texts = {
        "01-dpr-paper-en.pdf": pdf_text(ROOT / "01-dpr-paper-en.pdf"),
        "02-destatis-arbeitsmarkt-de.pdf": pdf_text(ROOT / "02-destatis-arbeitsmarkt-de.pdf"),
        "03-projekt-nordlicht-de.docx": docx_text(ROOT / "03-projekt-nordlicht-de.docx"),
        "04-grundgesetz-auszug-de.pdf": pdf_text(ROOT / "04-grundgesetz-auszug-de.pdf"),
        "05-nist-ai-rmf-scan-en.pdf": pdf_text(RAW / "nist.pdf", 5, 7),
        "06-bdsg-auszug-de.txt": (ROOT / "06-bdsg-auszug-de.txt").read_text(encoding="utf-8"),
        "07-wikipedia-rag-en.html": BeautifulSoup(
            (ROOT / "07-wikipedia-rag-en.html").read_text(encoding="utf-8"), "html.parser"
        ).get_text("\n"),
    }
    for name, text in texts.items():
        (OUT / f"{name}.txt").write_text(text, encoding="utf-8")


main()
